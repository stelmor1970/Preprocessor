"""docx 출력 (명세서 4절).

카드를 Word 문서로 출력한다. 답안용/암기용 뷰를 함께 싣고,
마크다운 비교표(| 항목 | 기준 |)는 실제 Word 표로 렌더링한다.

구조화 산출물이 필요한 경우 cards.json(Node docx 워크플로 연동용)을 사용한다.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

_TABLE_LINE = re.compile(r"^\s*\|.*\|\s*$")
_SEP_LINE = re.compile(r"^\s*\|[\s:\-|]+\|\s*$")


def _parse_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def _render_table(doc, block: list[str]) -> None:
    rows = [_parse_row(l) for l in block if not _SEP_LINE.match(l)]
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    for ri, r in enumerate(rows):
        cells = table.add_row().cells
        for j in range(ncols):
            val = r[j] if j < len(r) else ""
            cells[j].text = val
            if ri == 0:  # 헤더행 굵게
                for para in cells[j].paragraphs:
                    for run in para.runs:
                        run.bold = True


def _add_markdown_block(doc, text: str) -> None:
    """마크다운 텍스트를 문단/표로 렌더링."""
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        if _TABLE_LINE.match(lines[i]):
            block = []
            while i < len(lines) and _TABLE_LINE.match(lines[i]):
                block.append(lines[i])
                i += 1
            _render_table(doc, block)
        else:
            if lines[i].strip():
                doc.add_paragraph(lines[i])
            i += 1


def save_docx(cards, path: str | Path) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading("소방시설관리사 2차 실기 — 압축 카드", level=0)

    for c in cards:
        status = "채택" if c.adopted else "원문 유지(검증 실패)"
        doc.add_heading(f"[{c.id}]  · {status} · 압축률 {c.ratio:.0%}", level=1)

        if c.mnemonic:
            p = doc.add_paragraph()
            run = p.add_run(f"두음 암기어: {c.mnemonic}")
            run.bold = True

        doc.add_heading("📝 답안용 (시험 작성용)", level=2)
        _add_markdown_block(doc, c.compressed or "(없음)")

        if c.memo and c.memo != c.compressed:
            doc.add_heading("🧠 암기용 (약어 축약)", level=2)
            _add_markdown_block(doc, c.memo)

    doc.save(str(path))
